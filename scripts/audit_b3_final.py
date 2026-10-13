"""Single B3 package audit, original-Service replay and honest local-model probe."""
import ast
import base64
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
import contextlib
import copy
import hashlib
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import time
from urllib.request import Request,urlopen

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/"src"))
from scripts.build_b3_candidate import build_b3,BASELINE
from scripts.build_b_submission import load_rows,parse_submission,submission_sanity_check,publication_gate
from scripts.b2_experiments import fsh_checks,r4_structure
from scripts.audit_b_board import replay_service
from scripts.check_b_submission import CHILD

DEST=ROOT/"experiment/b3-final"
REPORTS=ROOT/"reports"


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path,data):
    Path(path).write_text(json.dumps(data,ensure_ascii=False,allow_nan=False,indent=2)+"\n",encoding="utf8")


def freeze():
    directory=DEST/"frozen";directory.mkdir(parents=True,exist_ok=True)
    prior=parse_submission((ROOT/"experiment/b2-official-contract/frozen/freeze_manifest.json").read_text())
    paths=[r["source"] for r in prior]+[
        "experiment/b2-official-contract/B2E.json","experiment/b2-official-contract/results.json",
        "submission/b2_candidate.manifest.json","src/b_board/b2_runtime.py","scripts/b2_experiments.py",
        "reports/B2_FINAL_VERDICT.md",
    ]
    records=[]
    for source in paths:
        target=directory/source
        target.parent.mkdir(parents=True,exist_ok=True)
        if not target.exists():
            shutil.copyfile(ROOT/source,target)
        if digest(target)!=digest(ROOT/source):
            raise ValueError("Frozen source changed: "+source)
        records.append({"source":source,"sha256":digest(target)})
    dump(directory/"manifest.json",records)
    return records


def source_audit(bundle):
    baseline=parse_submission(BASELINE.read_text(encoding="utf8"))
    result=[]
    for old,entry in zip(baseline["entry"][1:],bundle["entry"][1:]):
        r=entry["resource"]
        source=base64.b64decode(r["content"][0]["data"],validate=True).decode("utf8")
        previous=base64.b64decode(old["resource"]["content"][0]["data"]).decode("utf8")
        assert source.startswith(previous+"\n")
        tree=ast.parse(source)
        functions=[n for n in tree.body if isinstance(n,ast.FunctionDef)]
        names=Counter(n.name for n in functions)
        last={n.name:n for n in functions}
        for name,expected in (("evaluate","evaluate_b2"),("build_b_resources","build_b2_resources")):
            calls=[n.func.id for n in ast.walk(last[name]) if isinstance(n,ast.Call) and isinstance(n.func,ast.Name)]
            assert expected in calls
        namespace={};exec(compile(source,r["name"],"exec"),namespace)
        assert namespace["B2_SELECTED"]=="B2E" and namespace["B3_SELECTED"]=="B2E_WITH_RUNTIME_GUARDS"
        assert namespace["FHIRResourceBundleGenerator"]._transport is namespace["_b3_transport"]
        result.append({"profile_id":r["content"][0]["title"],"base64_utf8_compile":"PASS",
            "frozen_b2e_prefix":"PASS","final_evaluate":"evaluate_b2","final_builder":"build_b2_resources",
            "final_rescue":"B3 strict schema/literal evidence/event guard",
            "duplicate_definitions":{k:v for k,v in names.items() if v>1},
            "imports":sorted({n.module for n in ast.walk(tree) if isinstance(n,ast.ImportFrom)}),
            "socket_timeout_seconds":30,"max_tokens":1024,"source_sha256":hashlib.sha256(source.encode()).hexdigest()})
    return result


def load_generators(bundle):
    generators={}
    for entry in bundle["entry"][1:]:
        r=entry["resource"];namespace={}
        exec(compile(base64.b64decode(r["content"][0]["data"]),r["name"],"exec"),namespace)
        generators[r["content"][0]["title"]]=namespace
    return generators


def replay_one(case,namespace,row):
    generator=namespace["FHIRResourceBundleGenerator"]("http://localhost:3456")
    calls=[];generator.transport=lambda prompt:(calls.append(prompt) or "NO")
    started=time.perf_counter()
    with contextlib.redirect_stdout(io.StringIO()):
        response=generator.parse_clinical_text_to_fhir_bundle("p1",case["reports"])
    hit=replay_service(row,response)["hit"]
    return {"case":case["case"],"profile_id":case["profile_id"],"reports":case["reports"],
        "decision":generator.last_result["decision"],"service_hit":hit,"resources":response,
        "fsh_violations":fsh_checks(row,response),"r4_errors":r4_structure(response),
        "model_calls_stub":len(calls),"elapsed_seconds":time.perf_counter()-started}


def isolated(bundle,rows,cases):
    records=[]
    for row,entry in zip(rows,bundle["entry"][1:]):
        cid=str(row["profile_id"])
        case=next(c for c in cases if c["case"]==cid+"_positive_0")
        payload={"source":entry["resource"]["content"][0]["data"],"reports":case["reports"]}
        started=time.perf_counter()
        child=subprocess.run([sys.executable,"-I","-c",CHILD],input=json.dumps(payload),
            capture_output=True,text=True,timeout=30,cwd=DEST)
        if child.returncode:
            raise ValueError("Isolated failure "+cid+": "+child.stderr)
        response=parse_submission(child.stdout)
        if not response["entry"] or r4_structure(response) or not replay_service(row,response)["hit"]:
            raise ValueError("Isolated resource/service failure "+cid)
        records.append({"profile_id":cid,"positive_and_empty_execution":"PASS","original_service_hit":True,
            "elapsed_seconds":time.perf_counter()-started,"fsh_violations":fsh_checks(row,response)})
    return records


def publication_evidence(rows):
    import openpyxl
    evidence={"b_workbook":{},"other_workbooks":[],"headers":[],"historical_log":{}}
    path=Path("C:/Users/32433/Desktop/CHIP2026-CP2-data-test_b.xlsx")
    workbook=openpyxl.load_workbook(path,read_only=False,data_only=False)
    sheet=workbook["criterions"]
    header={cell.value:cell.column for cell in sheet[1]}
    cells=[]
    for number in (2,3):
        matching=next(i for i in range(2,sheet.max_row+1) if sheet.cell(i,header["profile_id"]).value==number)
        cells.append({"profile_id":number,"row":matching,"cells":{
            name:{"coordinate":sheet.cell(matching,header[name]).coordinate,"value":sheet.cell(matching,header[name]).value,
                  "data_type":sheet.cell(matching,header[name]).data_type} for name in ("id","index","deepquery_id")}})
    evidence["b_workbook"]={"path":str(path),"sha256":digest(path),"sheet":"criterions","raw_cells":cells}
    workbook.close()
    seen=set()
    for path in sorted(ROOT.parent.glob("CHIP2026*/data/*.xlsx")):
        sha=digest(path)
        if sha in seen:
            continue
        seen.add(sha)
        w=openpyxl.load_workbook(path,read_only=True,data_only=True)
        records=[]
        for sheet in w:
            values=iter(sheet.values);keys=next(values,())
            if "profile_id" not in keys:
                continue
            for value in values:
                record=dict(zip(keys,value))
                if str(record.get("profile_id")) in ("2","3"):
                    records.append({key:record.get(key) for key in ("profile_id","id","index","deepquery_id")})
        w.close()
        evidence["other_workbooks"].append({"path":str(path),"sha256":sha,"matching_2_3":records})
    header_paths=[ROOT.parent/"CHIP2026_CP2_A_baseline/data/train_set_message_bundle.json",
                  ROOT.parent/"CHIP2026_CP2_A_baseline_v2_4_3/submission/a_test_message_bundle_v2_4_3.json",
                  ROOT/"submission/a_test_message_bundle_v5s2_expe_aggressive_candidate.json",BASELINE]
    for path in header_paths:
        if path.exists():
            b=parse_submission(path.read_text(encoding="utf8"))
            evidence["headers"].append({"path":str(path),"sha256":digest(path),"message_header":b["entry"][0]})
    log=ROOT.parent/"CHIP2026_CP2_A_baseline_v2_4_3/evidence/online_a1_run.log"
    if log.exists():
        lines=log.read_text(encoding="utf8",errors="replace").splitlines()
        evidence["historical_log"]={"path":str(log),"sha256":digest(log),"routing_markers":{
            token:[i+1 for i,line in enumerate(lines) if token in line][:15]
            for token in ("train_set","test_a","test_b","chip2026","library-publish")}}
    evidence["search_scope"]=["current repo README/docs/reports/scripts/src/configs",
        "sibling CHIP2026 baseline README/config/startup files and unique data workbooks",
        "Desktop/Downloads and D:/browser-download filename inventory",
        "Codex attachments/cache: exact profile 2/3 deepquery-id search in text/json/html",
        "Codex browser cache session metadata inventory; no contest release document discovered"]
    evidence["search_limits"]=["unrelated milestone pytest temp directory denied access and excluded",
        "Windows process command-line enumeration denied; normal process inventory found no identified model server",
        "no live official announcement/checker source supplied or discovered; filesystem evidence cannot establish server routing"]
    return evidence


def live_probe(namespaces):
    result={"endpoint":"http://127.0.0.1:1213/v1/chat/completions","startup":"NO_PROJECT_DEFINED_TRUSTED_STARTUP_FOUND",
        "downloads":False,"external_paid_calls":False,"patient_data_sent":False,"tests":[]}
    started=time.perf_counter()
    try:
        with urlopen("http://127.0.0.1:1213/v1/models",timeout=5) as response:
            models=json.loads(response.read().decode())
        result["models_response"]=models
    except Exception as exc:
        result.update(status="UNAVAILABLE",availability_error=type(exc).__name__+": "+str(exc),
                      availability_seconds=time.perf_counter()-started)
        for name in ("NO_MATCH","POSITIVE","NEGATION_UNCERTAINTY","PARTIAL_JSON_RESCUE"):
            result["tests"].append({"case":name,"status":"NOT_RUN_ENDPOINT_UNAVAILABLE","real_model":False})
        result["fault_tests"]="Controlled malformed JSON and timeout tests are offline tests, not real model results"
        return result
    scenarios=[("NO_MATCH","病理阴性，未诊断淋巴瘤"),("POSITIVE","病理确诊DLBCL"),
        ("NEGATION_UNCERTAINTY","否认淋巴瘤病史，目前疑似淋巴瘤，病理待定"),
        ("PARTIAL_JSON_RESCUE","病理报告。确诊弥漫性大细胞B淋巴瘤")]
    for name,text in scenarios:
        reports=[{"text":text,"timestamp":"2026-10-09","event_id":"synthetic"}]
        prompt="只按原文提取事实，否定NEGATED，疑似UNCERTAIN，不补造事实。仅返回JSON facts，允许原子dlbcl和pathology，每项只有state/evidence/report。evidence逐字引用原文。病历："+json.dumps(reports,ensure_ascii=False)
        payload={"model":"local-model","temperature":0,"max_tokens":1024,"messages":[{"role":"user","content":prompt}]}
        record={"case":name,"synthetic_input":reports,"real_model":False};started=time.perf_counter()
        try:
            request=Request(result["endpoint"],json.dumps(payload).encode(),{"Content-Type":"application/json"})
            with urlopen(request,timeout=30) as response:
                body=json.loads(response.read().decode())
            record.update(real_model=True,response=body)
            content=body["choices"][0]["message"]["content"]
            record["parsed"]=namespaces["495"]["_b3_checked_facts"](content,"495",reports)
            record["schema_and_literal_evidence"]="PASS"
        except Exception as exc:
            record["error"]=type(exc).__name__+": "+str(exc)
        record["elapsed_seconds"]=time.perf_counter()-started;result["tests"].append(record)
    result["status"]="AVAILABLE" if all(r.get("schema_and_literal_evidence")=="PASS" for r in result["tests"]) else "AVAILABLE_WITH_FAILURES"
    result["fault_tests"]="Controlled malformed JSON and timeout tests are offline tests, not real model results"
    return result


def write_reports(data,baseline_cases,rows):
    publication=["# B3 Publication Contract","","PUBLICATION_IDENTITY_STATUS = UNRESOLVED","B_BOARD_READY = NO","",
        "## Traceable Original Evidence","",json.dumps(data["publication_evidence"],ensure_ascii=False,indent=2),"",
        "Original B cells contain literal id=tbd and empty index cells for profiles 2/3. This originates in the supplied workbook, not a B1 conversion. Distinct deepquery IDs identify criterion definitions but do not establish publishable Library identifiers. Unique historical A/train workbook copies contain neither B profile 2 nor 3. FSH describes clinical resources; Service queries patient resources and does not define Library routing.","",
        "The original training MessageHeader, historical A artifact and A Champion all carry source.challenge_id=train_set. Historical A run evidence establishes A subset routing; it contains no verified B mapping. B3 preserves the byte-identical B2E header for offline review. There is no evidence authorizing replacement with test_b or chip2026. Its B validity remains unresolved. The custom source.challenge_id is a platform field, not a generic R4 MessageHeader field; local package sanity is not a full R4 message validation.","",
        "No new identifier is guessed. cnwqk2/cnwqk3 and profile-2/profile-3 in the review package are inherited unverified B1/B2E placeholders. They are not official accepted mappings. Publication gate remains blocked on the actual serialized file.","",
        "## External Dependencies","",
        "Obtain organizer-issued profile 2/3 Library IDs/indexes and routing key; confirm the B MessageHeader challenge_id and required Library count; obtain official checker/IG or equivalent official validation. Save original source/date and validate the exact candidate hash. No automatic upload is performed, and local code changes cannot resolve these missing official facts."]
    publication[7]="```json\n"+publication[7]+"\n```"
    (REPORTS/"B3_PUBLICATION_CONTRACT.md").write_text("\n".join(publication)+"\n",encoding="utf8")
    violations=[r for r in baseline_cases if r["fsh_violations"]]
    fsh=["# B3 Remaining FSH Cases","","The 18 original B2E cases are enumerated below with exact inputs and emitted resources. Cases with multiple constraints count once. B3 retains 18 cases with existing local FSH subset violations. Full SUSHI/IG/profile validation remains unavailable.","",
         "12 cases: 85/205 fixed status=planned, which is not an R4 ProcedureStatus. 5 cases: 295 requires two components but evidence supports one. 1 case: 225 ongoing in-progress conflicts with completed and lacks an evidenced performed date. No 625 case has a violation detected by the existing subset checker; its clinical/Service logic disagreement remains separate.","",
         "No remaining case can be made fully compliant by a verified constructor-only change without changing evidence or the official constraints. For 225, adding the report timestamp as a completed-treatment time is unsupported. For 295, missing intent/quantity cannot be synthesized. For 85/205, using planned is invalid R4 and using completed for planned procedures falsifies status."]
    for record in violations:
        row=next(r for r in rows if str(r["profile_id"])==record["profile_id"])
        constraints=[line for line in row["FHIR_FSH"].splitlines() if line.startswith("* ") and any(word in line for word in ("status =","component ","component[","performedDateTime"))]
        fsh += ["","## "+record["case"],"","Profile: "+record["profile_id"]+"; known contract conflict: YES; safely fully repairable: NO.","",
            "Input evidence:","```json",json.dumps(record["reports"],ensure_ascii=False,indent=2),"```","",
            "Actual B2E FHIR:","```json",json.dumps(record["resources"],ensure_ascii=False,indent=2),"```","",
            "Original FSH constraints:","```fsh","\n".join(constraints),"```","",
            "Detected violations:",*['- '+v for v in record["fsh_violations"]]]
    (REPORTS/"B3_FSH_REMAINING.md").write_text("\n".join(fsh)+"\n",encoding="utf8")
    live=["# B3 Live Model Diagnosis","","MODEL_ENDPOINT_STATUS = "+data["live_model"]["status"],"",
        "Project and sibling configuration/startup scans found endpoint consumers and historical runtime-contract documents, but no defined trusted model-service launcher, model path, weights or deployment configuration for port 1213. No service was invented or downloaded. Normal process inventory identified no model server; privileged process command-line inventory was denied. This does not claim model software is uninstalled.","",
        "Actual localhost model-list request and status:","```json",json.dumps(data["live_model"],ensure_ascii=False,indent=2),"```","",
        "Reproduce availability with the bundled Python: urllib.request.urlopen('http://127.0.0.1:1213/v1/models', timeout=5). When a trusted project service is supplied, rerun this audit; the four synthetic real probes run only after availability succeeds. No genuine patient text is used.","",
        "Offline fault-injection tests verify malformed/duplicate/nonfinite JSON, extra fields, nonliteral evidence and TimeoutError degradation. They are explicitly not live-model tests. Production request timeout remains 30 seconds (socket timeout, not a guaranteed end-to-end deadline), no retry, max_tokens=1024. B3 records transport errors/duration instead of silently concealing failure. Token adequacy and real multi-patient model throughput remain unverified."]
    (REPORTS/"B3_LIVE_MODEL.md").write_text("\n".join(live)+"\n",encoding="utf8")
    verdict=["# B3 Final Verdict","","```text","A_CHAMPION = V5S.2 Experiment E (unchanged)",
        "B_OFFLINE_BEST = B3 (B2E plus reproduced runtime guards)","B_ONLINE_VERIFIED = NO","B_BOARD_READY = NO",
        "FSH_VIOLATIONS_REMAINING = "+str(data["summary"]["fsh_violating_cases"])+" cases (existing subset validator)",
        "MODEL_ENDPOINT_STATUS = "+data["live_model"]["status"],"PUBLICATION_IDENTITY_STATUS = UNRESOLVED","```","",
        "The manual-review candidate is experiment/b3-final/B3.json. Build/compile/reload checks pass, but it is not cleared for publication and is not an online champion. No Tianchi submission was performed.","",
        "## Verified Changes","",
        "B3 appends guards to the complete frozen B2E source. Three tests first reproduced: extra model JSON fields were accepted; categorical rescue could merge separate events; the actual embedded entrypoint retained pathology provenance from another event. B3 validates a closed facts schema, rejects duplicate/nonfinite JSON and requires literal source quotes, then recomputes facts in one linked event using the same single model response. Event guards preserve a genuine eligible event and cannot create numerical facts or override deterministic vetoes. Local transport now exposes its error/duration while preserving endpoint, token budget, timeout and deterministic NO fallback.","",
        "## Actual Package Evidence","","```json",json.dumps(data["summary"],ensure_ascii=False,indent=2),"```","",
        "All 16 final wrappers reach evaluate_b2 and build_b2_resources; frozen _B1 aliases intentionally provide underlying extraction/building, not a fallback to a previous submitted package. Duplicate definitions are recorded in results.json and their final bindings verified. Each final source keeps the exact B2E prefix; no later packaging/template pass overwrites it. 16/16 Libraries pass Python -I positive/empty execution and original unchanged offline Service queries. Concurrent entrypoint checks use distinct generator instances; last_result is per instance and a shared instance is not guaranteed to retain per-request diagnostics.","",
        "Strict serialized-file reload and local R4 subset pass; known FSH conflicts remain. Original Service code is replayed with offline FHIR mocks/frozen date, not a live official FHIR deployment. NO model stubs are used only for offline deterministic replay and counted. No gold labels: no precision/recall/F1 claim. The B2 six-strategy experiment was not repeated.","",
        "## External Dependencies and Stop Condition","",
        "Official profile 2/3 mapping, B MessageHeader contract and official checker/IG remain external dependencies. The 18 FSH cases are official specification/evidence conflicts, not a reason to invent values. Model startup requires a trusted project-defined service configuration. All available local checks have a finite stopping point; further ungrounded rule changes or blind validation uploads are not recommended."]
    (REPORTS/"B3_FINAL_VERDICT.md").write_text("\n".join(verdict)+"\n",encoding="utf8")


def run():
    DEST.mkdir(parents=True,exist_ok=True)
    frozen=freeze()
    old=parse_submission((ROOT/"experiment/b2-official-contract/results.json").read_text(encoding="utf8"))
    baseline_cases=[r for r in old["matrix"] if r["variant"]=="B2E"]
    rows=load_rows();row_by_id={str(r["profile_id"]):r for r in rows}
    bundle=build_b3();sources=source_audit(bundle)
    dump(DEST/"B3.json",bundle)
    actual=parse_submission((DEST/"B3.json").read_text(encoding="utf8"))
    package=submission_sanity_check(actual,rows)
    namespaces=load_generators(actual)
    matrix=[replay_one(c,namespaces[c["profile_id"]],row_by_id[c["profile_id"]]) for c in baseline_cases]
    isolated_results=isolated(actual,rows,baseline_cases)
    deltas=[{"case":r["case"],"before":b["decision"],"after":r["decision"]} for b,r in zip(baseline_cases,matrix) if b["decision"]!=r["decision"]]
    # Each worker owns a generator; concurrent calls do not share last_result state.
    tasks=[]
    for entry in actual["entry"][1:]:
        cid=entry["resource"]["content"][0]["title"]
        source=entry["resource"]["content"][0]["data"]
        case=next(c for c in baseline_cases if c["case"]==cid+"_positive_0")
        tasks.append((cid,source,case))
    def concurrent(task):
        cid,source,case=task
        payload={"source":source,"reports":case["reports"]}
        child=subprocess.run([sys.executable,"-I","-c",CHILD],input=json.dumps(payload),capture_output=True,text=True,timeout=30,cwd=DEST)
        return {"profile_id":cid,"returncode":child.returncode,
            "positive_nonempty":bool(parse_submission(child.stdout)["entry"]) if not child.returncode else False}
    with ThreadPoolExecutor(max_workers=4) as executor:
        concurrent_results=list(executor.map(concurrent,tasks))
    assert all(r["returncode"]==0 and r["positive_nonempty"] for r in concurrent_results)
    publication="UNEXPECTEDLY_PASSED"
    try:
        publication_gate(actual,rows)
    except ValueError as exc:
        publication="BLOCKED: "+str(exc)
    assert publication.startswith("BLOCKED")
    for raw in ('{"id":1,"id":2}','{"value":NaN}','{"value":Infinity}'):
        try:
            parse_submission(raw)
        except ValueError:
            continue
        raise ValueError("Strict JSON rejection failed")
    assert not any(r["r4_errors"] for r in matrix)
    assert not deltas
    summary={"synthetic_cases":len(matrix),"matches":sum(r["decision"]=="MATCH" for r in matrix),
        "service_hits":sum(r["service_hit"] for r in matrix),
        "service_differences":sum(r["service_hit"]!=(r["decision"]=="MATCH") for r in matrix),
        "b2e_deterministic_decision_deltas":len(deltas),"fsh_violating_cases":sum(bool(r["fsh_violations"]) for r in matrix),
        "local_r4_subset_violations":sum(bool(r["r4_errors"]) for r in matrix),
        "isolated_libraries":len(isolated_results),"concurrent_processes_passed":len(concurrent_results),
        "concurrent_workers":4,"offline_model_stub_calls":sum(r["model_calls_stub"] for r in matrix),
        "matrix_seconds":sum(r["elapsed_seconds"] for r in matrix),"publication_gate":publication,
        "formal_profile_validator":"UNAVAILABLE","gold_labels":"UNAVAILABLE"}
    data={"summary":summary,"package":package,"source_audit":sources,"matrix":matrix,"deltas":deltas,
          "isolated":isolated_results,"concurrency":concurrent_results,"frozen":frozen,
          "publication_evidence":publication_evidence(rows),"live_model":live_probe(namespaces)}
    for record in frozen:
        assert digest(ROOT/record["source"])==record["sha256"]
    dump(DEST/"results.json",data)
    dump(DEST/"B3.manifest.json",{"artifact":str(DEST/"B3.json"),"sha256":digest(DEST/"B3.json"),
        "base_b2e_sha256":digest(BASELINE),"publication_ready":False,"identity":"UNRESOLVED",
        "header_policy":"unchanged frozen B2E; B-specific contract unverified","frozen":frozen,"summary":summary})
    write_reports(data,baseline_cases,rows)
    print(json.dumps({"summary":summary,"model":data["live_model"]["status"],"artifact_sha256":digest(DEST/"B3.json")},indent=2))


if __name__=="__main__":
    run()
